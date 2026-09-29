# 記録表①②の見本から「1日ぶんのひな形」を取り出して JSON にする。
#
# 様式は罫線・結合・フォントが細かく、書き起こすと必ずずれる。
# 見本をそのまま型として持ち、日付と氏名だけ差し替える作りにする。
#
#   python extract-template.py <xlsx> <開始行> <行数> <列数> <出力json>
import sys, json
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L

src, top, nrows, ncols, out = sys.argv[1], int(sys.argv[2]), int(sys.argv[3]), int(sys.argv[4]), sys.argv[5]
wb = load_workbook(src)
ws = wb[wb.sheetnames[0]]


def argb(c):
    """openpyxl の色 → exceljs の ARGB。テーマ色は拾えないので None を返す。"""
    if c is None:
        return None
    if getattr(c, "type", None) == "rgb" and c.rgb and isinstance(c.rgb, str):
        return c.rgb if len(c.rgb) == 8 else ("FF" + c.rgb[-6:])
    return None


def side(s):
    if s is None or s.style is None:
        return None
    d = {"style": s.style}
    a = argb(s.color)
    d["color"] = {"argb": a or "FF000000"}
    return d


def style_of(cell):
    f, al, bd, fl = cell.font, cell.alignment, cell.border, cell.fill
    st = {}
    fo = {}
    if f.name: fo["name"] = f.name
    if f.size: fo["size"] = float(f.size)
    if f.bold: fo["bold"] = True
    if f.italic: fo["italic"] = True
    ca = argb(f.color)
    if ca and ca != "FF000000": fo["color"] = {"argb": ca}
    if fo: st["font"] = fo

    a = {}
    if al.horizontal: a["horizontal"] = al.horizontal
    if al.vertical: a["vertical"] = al.vertical
    if al.wrap_text: a["wrapText"] = True
    if al.shrink_to_fit: a["shrinkToFit"] = True
    if al.text_rotation: a["textRotation"] = al.text_rotation
    if al.indent: a["indent"] = int(al.indent)
    if a: st["alignment"] = a

    b = {}
    for k in ("top", "left", "bottom", "right"):
        s = side(getattr(bd, k))
        if s: b[k] = s
    if b: st["border"] = b

    if fl is not None and fl.fill_type == "solid":
        fg = argb(fl.fgColor)
        if fg and fg != "FFFFFFFF" and fg != "00000000":
            st["fill"] = {"type": "pattern", "pattern": "solid", "fgColor": {"argb": fg}}

    if cell.number_format and cell.number_format != "General":
        st["numFmt"] = cell.number_format
    return st


# スタイルは同じものが大量にあるので、表にまとめて番号で持つ
styles, index = [], {}


def style_id(st):
    key = json.dumps(st, sort_keys=True, ensure_ascii=False)
    if key not in index:
        index[key] = len(styles)
        styles.append(st)
    return index[key]


cells = []
for r in range(top, top + nrows):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=r, column=c)
        st = style_of(cell)
        v = cell.value
        if v is None and not st:
            continue
        item = {"r": r - top + 1, "c": c}
        if v is not None:
            item["v"] = v if isinstance(v, (int, float)) else str(v)
        if st:
            item["s"] = style_id(st)
        cells.append(item)

merges = []
for m in ws.merged_cells.ranges:
    if m.min_row >= top and m.max_row < top + nrows and m.max_col <= ncols:
        merges.append([m.min_row - top + 1, m.min_col, m.max_row - top + 1, m.max_col])

rowh = {}
for r in range(top, top + nrows):
    h = ws.row_dimensions[r].height
    if h:
        rowh[str(r - top + 1)] = float(h)

colw = {}
for c in range(1, ncols + 1):
    d = ws.column_dimensions.get(L(c))
    if d and d.width:
        colw[str(c)] = float(d.width)

ps = ws.page_setup
page = {
    "orientation": ps.orientation or "portrait",
    "paperSize": int(ps.paperSize) if ps.paperSize else 9,
    "fitToWidth": int(ps.fitToWidth) if ps.fitToWidth else 1,
    "fitToHeight": int(ps.fitToHeight) if ps.fitToHeight is not None else 0,
    "margins": {
        "left": float(ws.page_margins.left), "right": float(ws.page_margins.right),
        "top": float(ws.page_margins.top), "bottom": float(ws.page_margins.bottom),
        "header": float(ws.page_margins.header), "footer": float(ws.page_margins.footer),
    },
}

data = {"rows": nrows, "cols": ncols, "rowHeights": rowh, "colWidths": colw,
        "merges": merges, "styles": styles, "cells": cells, "page": page}
with open(out, "w", encoding="utf-8") as fp:
    json.dump(data, fp, ensure_ascii=False, separators=(",", ":"))

print(f"{out}  セル{len(cells)}個 / スタイル{len(styles)}種 / 結合{len(merges)}個 / {nrows}行x{ncols}列")
